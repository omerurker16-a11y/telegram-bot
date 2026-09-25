import io
import re
from collections import Counter
import logging
import pdfplumber
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from telegram import Update
from telegram.ext import ApplicationBuilder, ContextTypes, CommandHandler, MessageHandler, filters
import os

# Bot Token'ı (Render'da çevre değişkeninden alacak)
TOKEN = os.environ.get("BOT_TOKEN", "8963450215:AAGpTot613JEtzVMN2bjdHg1MC-xCPRvt40")

logging.basicConfig(format='%(asctime)s - %(name)s - %(levelname)s - %(message)s', level=logging.INFO)

def clean_ref(ref):
    """OCR hatalarını ve yapısal bozuklukları temizler."""
    if not ref:
        return ""
    r = str(ref).strip().replace('Θ', '0').replace('Ε', 'E').replace('Ο', '0')
    if len(r) >= 8 and r[-1] in ['8', '5']:
        r = r[:-1] + 'E'
    return r

def extract_data_from_pdf(pdf_bytes):
    """PDF içindeki tabloları/metinleri okur, Koltuk 1 ve Koltuk 2 referanslarını çıkarır."""
    parsed_data = []
    
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if not text:
                continue
            
            lines = text.split('\n')
            kat_bilgisi = "KAT 1"
            
            for line in lines:
                if "KAT 2" in line:
                    kat_bilgisi = "KAT 2"
                
                words = line.split()
                refs = [w for w in words if len(w) >= 8 and any(c.isdigit() for c in w)]
                
                if len(refs) >= 2:
                    try:
                        kit_ref = refs[0]
                        koltuk1 = clean_ref(refs[1]) if len(refs) > 1 else ""
                        koltuk2 = clean_ref(refs[2]) if len(refs) > 2 and refs[2].startswith(('73', '76')) else ""
                        
                        sira_no = words[0] if words[0].isdigit() else (words[1] if len(words)>1 and words[1].isdigit() else "-")
                        
                        parsed_data.append((sira_no, kat_bilgisi, koltuk1, koltuk2))
                    except Exception as e:
                        continue
                        
    return parsed_data

def create_excel(parsed_data, truck_name):
    """Verilerden kurallara uygun Excel dosyası üretir."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Yükleme Özeti"
    ws.views.sheetView[0].showGridLines = True
    
    k1_counts = Counter([clean_ref(k1) for _, _, k1, _ in parsed_data if clean_ref(k1)])
    k2_counts = Counter([clean_ref(k2) for _, _, _, k2 in parsed_data if clean_ref(k2)])

    header_fill_k1 = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_fill_k2 = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    title_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    
    font_title = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_regular = Font(name="Calibri", size=11)
    
    thin_border_side = Side(border_style="thin", color="D3D3D3")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("A1:G1")
    ws["A1"] = "PROSRS YÜKLEME LİSTESİ ÖZET ANALİZİ"
    ws["A1"].font = font_title
    ws["A1"].fill = title_fill
    ws["A1"].alignment = align_center
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A4:C4")
    ws["A4"] = "KOLTUK 1 FREKANS ÖZETİ"
    ws["A4"].font = font_header
    ws["A4"].fill = header_fill_k1
    ws["A4"].alignment = align_center
    
    ws.merge_cells("E4:G4")
    ws["E4"] = "KOLTUK 2 FREKANS ÖZETİ"
    ws["E4"].font = font_header
    ws["E4"].fill = header_fill_k2
    ws["E4"].alignment = align_center

    for col, text in zip(["A", "B", "C"], ["Sıra", "Referans Kodu", "Adet"]):
        ws[f"{col}5"] = text
        ws[f"{col}5"].font = font_header
        ws[f"{col}5"].fill = header_fill_k1
        ws[f"{col}5"].border = thin_border
        
    for col, text in zip(["E", "F", "G"], ["Sıra", "Referans Kodu", "Adet"]):
        ws[f"{col}5"] = text
        ws[f"{col}5"].font = font_header
        ws[f"{col}5"].fill = header_fill_k2
        ws[f"{col}5"].border = thin_border

    r1 = 6
    for idx, (ref, count) in enumerate(k1_counts.most_common(), 1):
        ws[f"A{r1}"] = idx
        ws[f"B{r1}"] = ref
        ws[f"C{r1}"] = count
        for c in ["A", "B", "C"]:
            ws[f"{c}{r1}"].border = thin_border
            if idx % 2 == 0: ws[f"{c}{r1}"].fill = zebra_fill
        r1 += 1
        
    ws[f"A{r1}"] = "TOPLAM"
    ws.merge_cells(f"A{r1}:B{r1}")
    ws[f"A{r1}"].font = font_bold
    ws[f"C{r1}"] = f"=SUM(C6:C{r1-1})"
    ws[f"C{r1}"].font = font_bold
    
    r2 = 6
    for idx, (ref, count) in enumerate(k2_counts.most_common(), 1):
        ws[f"E{r2}"] = idx
        ws[f"F{r2}"] = ref
        ws[f"G{r2}"] = count
        for c in ["E", "F", "G"]:
            ws[f"{c}{r2}"].border = thin_border
            if idx % 2 == 0: ws[f"{c}{r2}"].fill = zebra_fill
        r2 += 1

    ws[f"E{r2}"] = "TOPLAM"
    ws.merge_cells(f"E{r2}:F{r2}")
    ws[f"E{r2}"].font = font_bold
    ws[f"G{r2}"] = f"=SUM(G6:G{r2-1})"
    ws[f"G{r2}"].font = font_bold

    start_r = max(r1, r2) + 3
    ws.merge_cells(f"A{start_r}:D{start_r}")
    ws[f"A{start_r}"] = "TÜM PALET YÜKLEME LİSTESİ DÖKÜMÜ"
    ws[f"A{start_r}"].font = font_header
    ws[f"A{start_r}"].fill = header_fill_k1
    
    headers = ["Palet Sıra No", "Kat", "Koltuk 1 Referansı", "Koltuk 2 Referansı"]
    for c, h in zip(["A", "B", "C", "D"], headers):
        ws[f"{c}{start_r+1}"] = h
        ws[f"{c}{start_r+1}"].font = font_header
        ws[f"{c}{start_r+1}"].fill = header_fill_k2
        
    cur_r = start_r + 2
    for sira, kat, k1, k2 in parsed_data:
        ws[f"A{cur_r}"] = sira
        ws[f"B{cur_r}"] = kat
        ws[f"C{cur_r}"] = k1 if k1 else "-"
        ws[f"D{cur_r}"] = k2 if k2 else "-"
        for c in ["A", "B", "C", "D"]:
            ws[f"{c}{cur_r}"].border = thin_border
        cur_r += 1

    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["F"].width = 20
    
    excel_io = io.BytesIO()
    wb.save(excel_io)
    excel_io.seek(0)
    
    return excel_io, k1_counts, k2_counts

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "👋 Merhaba! Ben PROSRS Yükleme Asistanı.\n\n"
        "Bana PDF formatındaki yükleme listenizi gönderin, "
        "Koltuk 1 ve Koltuk 2 referanslarını saniyeler içinde temizleyip Excel tablosu olarak size göndereyim."
    )

async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    doc = update.message.document
    if not doc.file_name.lower().endswith('.pdf'):
        await update.message.reply_text("❌ Lütfen sadece PDF dosyası gönderin.")
        return

    msg = await update.message.reply_text("⏳ Yükleme listesi işleniyor, lütfen bekleyin...")
    
    try:
        file = await context.bot.get_file(doc.file_id)
        pdf_bytes = await file.download_as_bytearray()
        
        parsed_data = extract_data_from_pdf(pdf_bytes)
        
        if not parsed_data:
            await msg.edit_text("⚠️ PDF içinde veri bulunamadı veya format uygun değil.")
            return
            
        excel_io, k1_counts, k2_counts = create_excel(parsed_data, doc.file_name)
        
        top_k1 = k1_counts.most_common(1)[0] if k1_counts else ("Yok", 0)
        top_k2 = k2_counts.most_common(1)[0] if k2_counts else ("Yok", 0)
        
        ozet_metin = (
            "✅ **İşlem Tamamlandı!**\n\n"
            f"🥇 **Koltuk 1 En Çok Tekrarlayan:** {top_k1[0]} ({top_k1[1]} adet)\n"
            f"🥇 **Koltuk 2 En Çok Tekrarlayan:** {top_k2[0]} ({top_k2[1]} adet)\n\n"
            "Detaylı Excel raporunuz aşağıdadır 👇"
        )
        
        excel_filename = f"{doc.file_name.replace('.pdf', '')}_Analiz.xlsx"
        
        await msg.delete()
        await context.bot.send_document(
            chat_id=update.message.chat_id,
            document=excel_io,
            filename=excel_filename,
            caption=ozet_metin,
            parse_mode='Markdown'
        )
        
    except Exception as e:
        logging.error(f"Hata: {e}")
        await msg.edit_text(f"❌ İşlem sırasında bir hata oluştu: {str(e)}")

if __name__ == "__main__":
    from flask import Flask
    from threading import Thread

    if not TOKEN:
        raise SystemExit("BOT_TOKEN ortam degiskeni yok. Render > Environment'dan ekleyin.")

    web = Flask(__name__)

    @web.route("/")
    def health():
        return "Bot calisiyor", 200

    def run_web():
        port = int(os.environ.get("PORT", 10000))
        web.run(host="0.0.0.0", port=port)

    Thread(target=run_web, daemon=True).start()

    application = ApplicationBuilder().token(TOKEN).build()
    application.add_handler(CommandHandler("start", start))
    application.add_handler(MessageHandler(filters.Document.ALL, handle_document))
    print("Bot basariyla baslatildi ve 7/24 calisiyor!")
    application.run_polling(allowed_updates=Update.ALL_TYPES, drop_pending_updates=True)
